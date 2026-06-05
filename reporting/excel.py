"""Excel report export (openpyxl).

Three sheets:
  * Summary    — run metadata (reproducibility stamps) + headline KPIs + sign-off block.
  * By Field   — mismatch/exception breakdown per field (the "category" view).
  * Drilldown  — every field-level result row, with a PDF Ref column that traces the
                 rule back to its source clause (populated once P3 citations exist).

Every export embeds template_version / mapping / rule_engine_version / run_id so the
file is self-describing and reproducible.
"""
from __future__ import annotations
from typing import Any, Dict, List, Tuple
import os

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# palette
HEAD = PatternFill("solid", fgColor="1F2937")
HEAD_FONT = Font(color="FFFFFF", bold=True)
KPI_FONT = Font(bold=True, size=14)
CAT_FILL = {
    "MATCH": "DCFCE7", "MISMATCH": "FEE2E2", "EXCEPTION": "FEF3C7",
    "DUPLICATE": "E0E7FF", "MISSING_IN_SOURCE": "F3E8FF", "MISSING_IN_DEST": "F3E8FF",
}
THIN = Border(*([Side(style="thin", color="E5E7EB")] * 4))


def _header(ws, row: int, cols: List[str]) -> None:
    for i, c in enumerate(cols, start=1):
        cell = ws.cell(row=row, column=i, value=c)
        cell.fill = HEAD; cell.font = HEAD_FONT; cell.border = THIN
        cell.alignment = Alignment(vertical="center")


def _autosize(ws, max_width: int = 60) -> None:
    for col in ws.columns:
        width = max((len(str(c.value)) for c in col if c.value is not None), default=10)
        ws.column_dimensions[get_column_letter(col[0].column)].width = min(width + 2, max_width)


def _citation(detail: Any) -> str:
    if isinstance(detail, dict):
        cit = detail.get("citation")
        if isinstance(cit, dict):
            page = cit.get("page")
            txt = cit.get("source_text", "")
            return f"p.{page}: {txt}" if page else (txt or "—")
    return "—"   # populated when P3 PDF extraction provides citations


def build_report(run: Dict[str, Any], field_results: List[Dict[str, Any]],
                 rollup: List[Tuple[str, str]], out_path: str) -> str:
    wb = Workbook()

    # ---------- Sheet 1: Summary ----------
    ws = wb.active; ws.title = "Summary"
    ws["A1"] = "Reconciliation Report"; ws["A1"].font = Font(bold=True, size=16)

    meta = [
        ("Run ID", run.get("run_id")),
        ("Status", run.get("status")),
        ("Rule engine version", run.get("rule_engine_version")),
        ("As-of date", run.get("as_of_date")),
        ("Template version", run.get("template_version_id")),
        ("Mapping", run.get("mapping_id")),
        ("Source dataset", run.get("source_dataset_id")),
        ("Destination dataset", run.get("dest_dataset_id")),
        ("Finished at", run.get("finished_at")),
    ]
    r = 3
    for k, v in meta:
        ws.cell(row=r, column=1, value=k).font = Font(bold=True)
        ws.cell(row=r, column=2, value=v)
        r += 1

    s = run.get("summary") or {}
    r += 1
    ws.cell(row=r, column=1, value="Results").font = Font(bold=True, size=13); r += 1
    kpis = [
        ("Total", s.get("total_records", 0)),
        ("✅ Match", s.get("match_count", 0)),
        ("❌ Mismatch", s.get("mismatch_count", 0)),
        ("⚠ Missing (src)", s.get("missing_src", 0)),
        ("⚠ Missing (dst)", s.get("missing_dst", 0)),
        ("🔁 Duplicate", s.get("duplicate_count", 0)),
        ("🚫 Exception", s.get("exception_count", 0)),
    ]
    for i, (label, val) in enumerate(kpis):
        ws.cell(row=r, column=1 + i, value=label).font = Font(bold=True)
        c = ws.cell(row=r + 1, column=1 + i, value=val); c.font = KPI_FONT
    r += 3

    # sign-off block (audit)
    ws.cell(row=r, column=1, value="Sign-off").font = Font(bold=True, size=13); r += 1
    for label in ("Reviewed by", "Approved by", "Date"):
        ws.cell(row=r, column=1, value=label).font = Font(bold=True)
        ws.cell(row=r, column=2, value="________________________")
        r += 1
    _autosize(ws)

    # ---------- Sheet 2: By Field ----------
    ws2 = wb.create_sheet("By Field")
    _header(ws2, 1, ["Field", "Mismatch", "Exception", "Total issues"])
    agg: Dict[str, Dict[str, int]] = {}
    for row in field_results:
        if row["category"] in ("MISMATCH", "EXCEPTION"):
            d = agg.setdefault(row["field_key"], {"MISMATCH": 0, "EXCEPTION": 0})
            d[row["category"]] += 1
    rr = 2
    for fld, d in sorted(agg.items(), key=lambda kv: -(kv[1]["MISMATCH"] + kv[1]["EXCEPTION"])):
        ws2.cell(row=rr, column=1, value=fld)
        ws2.cell(row=rr, column=2, value=d["MISMATCH"])
        ws2.cell(row=rr, column=3, value=d["EXCEPTION"])
        ws2.cell(row=rr, column=4, value=d["MISMATCH"] + d["EXCEPTION"])
        rr += 1
    if not agg:
        ws2.cell(row=2, column=1, value="(no field-level issues)")
    _autosize(ws2)

    # ---------- Sheet 3: Drilldown ----------
    ws3 = wb.create_sheet("Drilldown")
    cols = ["Record Key", "Field", "Category", "Severity", "Rule", "Expected",
            "Actual", "Verdict", "PDF Ref"]
    _header(ws3, 1, cols)
    rr = 2
    for row in field_results:
        vals = [row["record_key"], row["field_key"], row["category"], row.get("severity"),
                row.get("rule_id"), row.get("expected"), row.get("actual"),
                row.get("verdict"), _citation(row.get("detail"))]
        for i, v in enumerate(vals, start=1):
            cell = ws3.cell(row=rr, column=i, value=v); cell.border = THIN
        fill = CAT_FILL.get(row["category"])
        if fill:
            ws3.cell(row=rr, column=3).fill = PatternFill("solid", fgColor=fill)
        rr += 1
    ws3.freeze_panes = "A2"
    ws3.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{max(rr-1,1)}"
    _autosize(ws3)

    os.makedirs(os.path.dirname(os.path.abspath(out_path)), exist_ok=True)
    wb.save(out_path)
    return out_path
